"""Export and validate a completed EHQ adjudication workbook as UTF-8 CSV."""

from __future__ import annotations

import argparse
import csv
import io
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from ehq.artifacts import atomic_write_text  # noqa: E402
from ehq.constants import RESPONSE_LABELS  # noqa: E402


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Workbook export requires the optional reporting dependency: "
            "pip install 'ehq[reporting]'"
        ) from exc

    if args.output.exists():
        raise FileExistsError(f"Output already exists: {args.output}")
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    if "Adjudication" not in workbook.sheetnames:
        raise ValueError("Workbook does not contain an Adjudication sheet")
    sheet = workbook["Adjudication"]
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        raise ValueError("Adjudication sheet is empty")
    headers = [str(value or "").strip() for value in values[0]]
    required = {"validation_id", "category", "human_label", "notes"}
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"Missing required columns: {missing}")
    index = {name: headers.index(name) for name in required}

    rows = []
    for values_row in values[1:]:
        item_id = str(values_row[index["validation_id"]] or "").strip()
        if not item_id:
            continue
        category = str(values_row[index["category"]] or "").strip()
        label = str(values_row[index["human_label"]] or "").strip()
        notes = str(values_row[index["notes"]] or "").strip()
        if label not in RESPONSE_LABELS:
            raise ValueError(f"Invalid human_label for {item_id}: {label!r}")
        if category in {"FEQ", "CCQ"} and label == "CONFIDENT_CORRECT":
            raise ValueError(
                f"CONFIDENT_CORRECT is prohibited for {category}: {item_id}"
            )
        if not notes:
            raise ValueError(f"Adjudication notes are required for {item_id}")
        rows.append(
            {
                "validation_id": item_id,
                "human_label": label,
                "notes": notes,
            }
        )

    if len(rows) != 25:
        raise ValueError(f"Expected 25 adjudications, found {len(rows)}")
    ids = [row["validation_id"] for row in rows]
    if len(set(ids)) != len(ids):
        raise ValueError("Duplicate validation_id in adjudication workbook")

    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(
        buffer,
        fieldnames=["validation_id", "human_label", "notes"],
    )
    writer.writeheader()
    writer.writerows(rows)
    atomic_write_text(args.output, "\ufeff" + buffer.getvalue())
    print(
        json.dumps(
            {
                "output": str(args.output.resolve()),
                "n_adjudications": len(rows),
                "status": "READY_FOR_FINAL_ANALYSIS",
            },
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
